"""
Knowledge Agent Tools
文件系统分析 + 文件内容提取工具集
支持PDF/Excel/图片转图片后VLM分析
"""

import os
import base64
import subprocess
import tempfile
from pathlib import Path
from typing import Dict, List, Optional, Tuple
from datetime import datetime

# PyMuPDF for PDF
try:
    import fitz
    HAS_FITZ = True
except ImportError:
    HAS_FITZ = False

# pdf2image for Excel->PDF->Image
try:
    from pdf2image import convert_from_path
    HAS_PDF2IMAGE = True
except ImportError:
    HAS_PDF2IMAGE = False

# 业务关键词映射
BUSINESS_KEYWORDS = {
    "Admin": "行政管理",
    "Corporate": "公司法人",
    "Finance": "财务税务",
    "HR": "人力资源",
    "Production": "生产制造",
    "Sales": "销售市场",
    "Logistics": "物流供应链",
    "Quality": "质量管理",
    "R&D": "研发技术",
    "Legal": "法务合规",
    "IT": "信息技术",
    "Hungary": "匈牙利项目",
    "Invoice": "发票账单",
    "Contract": "合同协议",
    "Report": "报告文档",
    "Certificate": "资质证书",
    "PO": "采购订单",
    "Quote": "报价单",
    "Packing": "装箱单",
}

# 文件类型分类
FILE_CATEGORIES = {
    "合同文档": [".pdf", ".doc", ".docx"],
    "财务数据": [".xlsx", ".xls", ".csv"],
    "技术图纸": [".dwg", ".dxf", ".cad"],
    "图片资料": [".png", ".jpg", ".jpeg", ".gif", ".bmp"],
    "压缩包": [".zip", ".rar", ".7z"],
    "演示文稿": [".ppt", ".pptx"],
}

# 可VLM分析的文件类型
VLM_SUPPORTED_TYPES = {
    "direct_image": [".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".tif"],
    "pdf": [".pdf"],
    "excel": [".xlsx", ".xls"],
    "word": [".doc", ".docx"],
}


def analyze_folder_structure(base_path: str, folder_path: str = "") -> Dict:
    """
    深度分析文件夹结构，提取关键信息
    """
    target_dir = os.path.join(base_path, folder_path) if folder_path else base_path

    if not os.path.exists(target_dir):
        return {"error": "文件夹不存在", "folder_path": folder_path}

    if not os.path.isdir(target_dir):
        return {"error": "路径不是文件夹", "folder_path": folder_path}

    folders = []
    files = []
    file_types = {}
    total_size = 0

    try:
        entries = os.listdir(target_dir)
    except PermissionError:
        return {"error": "没有访问权限", "folder_path": folder_path}

    for entry in entries:
        entry_path = os.path.join(target_dir, entry)
        rel_path = os.path.join(folder_path, entry) if folder_path else entry

        if os.path.isdir(entry_path):
            file_count = _count_files_recursive(entry_path)
            business_hint = classify_business_type(entry)
            folders.append({
                "name": entry,
                "path": rel_path,
                "file_count": file_count,
                "business_hint": business_hint
            })
        else:
            ext = os.path.splitext(entry)[1].lower()
            try:
                stat = os.stat(entry_path)
                size = stat.st_size
                mtime = datetime.fromtimestamp(stat.st_mtime)
            except OSError:
                size = 0
                mtime = None

            total_size += size
            file_types[ext] = file_types.get(ext, 0) + 1

            # 检查是否可VLM分析
            vlm_type = get_vlm_type(ext)

            files.append({
                "name": entry,
                "path": rel_path,
                "ext": ext,
                "size": size,
                "size_str": _format_size(size),
                "modified": mtime.isoformat() if mtime else None,
                "vlm_supported": vlm_type is not None,
                "vlm_type": vlm_type
            })

    folders.sort(key=lambda x: x["name"].lower())
    files.sort(key=lambda x: x["name"].lower())
    categorized = _categorize_files(file_types)
    depth = len(folder_path.split('/')) if folder_path else 0

    # 统计可VLM分析的文件数
    vlm_files = [f for f in files if f.get("vlm_supported")]

    return {
        "folder_path": folder_path or "根目录",
        "depth": depth,
        "folder_count": len(folders),
        "file_count": len(files),
        "vlm_file_count": len(vlm_files),
        "total_size": total_size,
        "total_size_str": _format_size(total_size),
        "folders": folders,
        "files": files[:30],
        "file_types": file_types,
        "categorized": categorized,
        "has_more_files": len(files) > 30
    }


def get_vlm_type(ext: str) -> Optional[str]:
    """获取文件的VLM处理类型"""
    ext = ext.lower()
    for vlm_type, extensions in VLM_SUPPORTED_TYPES.items():
        if ext in extensions:
            return vlm_type
    return None


def file_to_images(file_path: str, max_pages: int = 5, dpi: int = 150) -> List[Dict]:
    """
    将文件转换为图片列表（base64格式）

    支持:
    - 图片: 直接转base64
    - PDF: 用PyMuPDF逐页转图片
    - Excel: LibreOffice转PDF再转图片
    - Word: LibreOffice转PDF再转图片

    Returns:
        List[{"page": int, "base64": str, "mime": str}]
    """
    path = Path(file_path)
    ext = path.suffix.lower()
    vlm_type = get_vlm_type(ext)

    if vlm_type is None:
        return [{"error": f"不支持的文件类型: {ext}"}]

    if vlm_type == "direct_image":
        return _image_to_base64_list(path)
    elif vlm_type == "pdf":
        return _pdf_to_images(path, max_pages, dpi)
    elif vlm_type in ["excel", "word"]:
        return _office_to_images(path, max_pages, dpi)

    return [{"error": "未知处理方式"}]


def _image_to_base64_list(path: Path) -> List[Dict]:
    """图片直接转base64"""
    try:
        img_bytes = path.read_bytes()
        b64 = base64.b64encode(img_bytes).decode()
        mime = _get_mime_type(path)
        return [{
            "page": 1,
            "base64": b64,
            "mime": mime,
            "size": len(img_bytes)
        }]
    except Exception as e:
        return [{"error": str(e)}]


def _pdf_to_images(path: Path, max_pages: int = 5, dpi: int = 150) -> List[Dict]:
    """PDF转图片列表"""
    if not HAS_FITZ:
        return [{"error": "PyMuPDF (fitz) 未安装"}]

    try:
        doc = fitz.open(str(path))
        page_count = len(doc)
        results = []

        for page_num in range(min(page_count, max_pages)):
            page = doc[page_num]
            mat = fitz.Matrix(dpi/72, dpi/72)
            pix = page.get_pixmap(matrix=mat)
            img_bytes = pix.tobytes("png")
            b64 = base64.b64encode(img_bytes).decode()

            results.append({
                "page": page_num + 1,
                "base64": b64,
                "mime": "image/png",
                "size": len(img_bytes)
            })

        doc.close()

        if page_count > max_pages:
            results.append({"info": f"共{page_count}页，仅处理前{max_pages}页"})

        return results
    except Exception as e:
        return [{"error": str(e)}]


def _office_to_images(path: Path, max_pages: int = 5, dpi: int = 150) -> List[Dict]:
    """Office文档(Excel/Word)转图片"""
    if not HAS_PDF2IMAGE:
        return [{"error": "pdf2image 未安装"}]

    try:
        with tempfile.TemporaryDirectory() as tmpdir:
            # Step 1: LibreOffice 转 PDF
            result = subprocess.run([
                "libreoffice", "--headless", "--convert-to", "pdf",
                "--outdir", tmpdir, str(path)
            ], capture_output=True, timeout=120)

            if result.returncode != 0:
                stderr = result.stderr.decode()[:200] if result.stderr else "Unknown error"
                return [{"error": f"LibreOffice转换失败: {stderr}"}]

            # 找到生成的PDF
            pdf_files = list(Path(tmpdir).glob("*.pdf"))
            if not pdf_files:
                return [{"error": "LibreOffice未生成PDF"}]

            pdf_path = pdf_files[0]

            # Step 2: PDF转图片
            images = convert_from_path(str(pdf_path), dpi=dpi)
            results = []

            for i, img in enumerate(images[:max_pages]):
                with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
                    img.save(tmp.name, "PNG")
                    img_bytes = Path(tmp.name).read_bytes()
                    os.unlink(tmp.name)

                b64 = base64.b64encode(img_bytes).decode()
                results.append({
                    "page": i + 1,
                    "base64": b64,
                    "mime": "image/png",
                    "size": len(img_bytes)
                })

            if len(images) > max_pages:
                results.append({"info": f"共{len(images)}页，仅处理前{max_pages}页"})

            return results

    except subprocess.TimeoutExpired:
        return [{"error": "LibreOffice转换超时"}]
    except Exception as e:
        return [{"error": str(e)}]


def classify_business_type(name: str) -> Optional[str]:
    """根据名称分类业务类型"""
    name_lower = name.lower()
    for keyword, description in BUSINESS_KEYWORDS.items():
        if keyword.lower() in name_lower:
            return description
    return None


def _count_files_recursive(dir_path: str) -> int:
    """递归统计目录下所有文件数量"""
    total = 0
    try:
        for entry in os.listdir(dir_path):
            entry_path = os.path.join(dir_path, entry)
            if os.path.isfile(entry_path):
                total += 1
            elif os.path.isdir(entry_path):
                total += _count_files_recursive(entry_path)
    except PermissionError:
        pass
    return total


def _format_size(size: int) -> str:
    """格式化文件大小"""
    if size < 1024:
        return f"{size}B"
    elif size < 1024 * 1024:
        return f"{size // 1024}KB"
    elif size < 1024 * 1024 * 1024:
        return f"{size // (1024 * 1024)}MB"
    else:
        return f"{size // (1024 * 1024 * 1024)}GB"


def _categorize_files(file_types: Dict[str, int]) -> Dict[str, int]:
    """根据扩展名统计文件类别"""
    categorized = {}
    for category, extensions in FILE_CATEGORIES.items():
        count = sum(file_types.get(ext, 0) for ext in extensions)
        if count > 0:
            categorized[category] = count
    return categorized


def _get_mime_type(path: Path) -> str:
    """获取MIME类型"""
    suffix = path.suffix.lower()
    mime_map = {
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".gif": "image/gif",
        ".bmp": "image/bmp",
        ".tif": "image/tiff",
        ".tiff": "image/tiff",
        ".pdf": "application/pdf",
    }
    return mime_map.get(suffix, "application/octet-stream")


def extract_word_text(file_path: str) -> str:
    """
    从Word文档提取纯文本
    
    Args:
        file_path: Word文档路径 (.docx)
        
    Returns:
        提取的文本内容
    """
    try:
        from docx import Document
        doc = Document(file_path)
        
        paragraphs = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)
        
        # 也提取表格内容
        for table in doc.tables:
            for row in table.rows:
                row_text = ' | '.join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        
        return '\n'.join(paragraphs)
    except Exception as e:
        return f"提取Word文本失败: {str(e)}"
